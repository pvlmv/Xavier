import re
import csv
import os
from datetime import datetime
from docx import Document
from PySide6.QtCore import Slot

@Slot()
def generate_docx_with_applied_item(template_path: str, output_path: str, item : dict[str, str | list[str]]):
    Document(template_path).save(output_path)
    document = Document(output_path)
    values = dict(item)
    
    values["current_date_formated(%d.%m.%Y)"] = datetime.now().strftime("%d.%m.%Y") # TO BE MODIFIED: Add dynamic date formating
    
    for paragraph in document.paragraphs:
        for variable, value in values.items():
            paragraph.text = re.sub(rf"\$`{re.escape(variable)}`", str(value), paragraph.text)
    document.save(output_path)
    # os.system(f"start {output_path}") Open generated document
    
@Slot()
def get_item_list_from_csv(csv_path: str) -> list[dict[str, str]]:
    with open(csv_path, 'r') as f:
        item_list = [{k: v for k, v in row.items()} for row in csv.DictReader(f, skipinitialspace=True)]
    return item_list

@Slot()
def expect_items_from_docx(template_path: str) -> dict[str, list[str]]:
    document = Document(template_path)
    text = "\n".join([paragraph.text for paragraph in document.paragraphs])
    unit_variables = [re.sub("\\$?\\`","",var) for var in list(set(re.findall("\\$`\\w+`", text)))]
    item_variables = [re.sub("\\$?\\`","",var) for var in list(set(re.findall("\\$`\\w+[.]\\w+`", text)))]

    items = dict()
    for item_name in list(set([var.split('.')[0] for var in item_variables])):
        items[item_name] = list(set([var.split('.')[1] for var in item_variables if var.split('.')[0]==item_name]))
    items["UNIT"] = unit_variables
    for item_var_lists in items.values(): item_var_lists.sort()
    
    return items

@Slot()
def generate_final_input_item(expected_attributes:dict[str, str | list[str]], chosen_attributes:dict[str, str | list[str]]) -> dict[str, str]:
    final = dict()
    
    for item_name in [item_name for item_name in expected_attributes.keys() if item_name != "UNIT"]:
        item = [item for item in get_item_list_from_csv(chosen_attributes[item_name]['path']) if item['id']==chosen_attributes[item_name]['id']][0]
        for var in item.keys():
            final[f"{item_name}.{var}"] = item[var]
            
    for unit_name in expected_attributes['UNIT']:
        final[unit_name] = chosen_attributes[unit_name]

    return final


@Slot()
def find_fitting_items(expected_attributes : list[str]) -> dict[str, str]:
    PATH = "./items"
    item_files = os.listdir(PATH)
    fitting_items = dict()
    for item_file in item_files:
        items = get_item_list_from_csv(f"{PATH}/{item_file}")
        if all(items[0].__contains__(x) for x in expected_attributes):
            fitting_items.update({item["id"]:item_file for item in items})
                
    return fitting_items