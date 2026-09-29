import re
import csv
import os
from datetime import datetime
from docx import Document
from PySide6.QtCore import Slot

def substitute_with_current_date(text:str) -> str:
    def get_date_format(date:str="") -> str:
        format = "%d.%m.%Y"
        if date.__contains__("(") and date.replace("(","")[1].split(")")[0].strip()!="":
            format = date.split("(")[1].split(")")[0]
        else:
            try:
                with open("./contents/config.txt") as f:
                    for line in f.readlines():
                        if line.split(' ')[0] == "default_date_format": 
                            format = " ".join(line.split(' ')[1::]).strip()
                            break
                        else: format = "%d.%m.%Y"
            except:
                format = "%d.%m.%Y"
        return format
    
    text = re.sub(rf"\$`current_date_formatted\(\)`|\$`current_date_formatted`",datetime.now().strftime(get_date_format()),text)
    
    date_to_replace = re.findall(rf"\$`current_date_formatted\(?.*\)?`", text)
    if date_to_replace == []:
        return text
    
    for date in date_to_replace:
        format = get_date_format(date)
        text = re.sub(rf"\$`current_date_formatted\({re.escape(format)}\)`", datetime.now().strftime(format), text)
    return text

PRESET_VARIABLES = {
    "current_date_formatted": substitute_with_current_date,
}

def open_word_document(docx_path: str):
    if os.name != "nt":
        raise OSError("Opening a Word document with the Windows file association requires Windows.")

    absolute_path = os.path.abspath(docx_path)
    if not os.path.isfile(absolute_path):
        raise FileNotFoundError(f"Word document not found: {absolute_path}")

    os.startfile(absolute_path, "open")


@Slot()
def generate_docx_with_applied_item(template_path: str, output_path: str, item : dict[str, str | list[str]]):
    Document(template_path).save(output_path)
    document = Document(output_path)
    values : dict[str, callable] = {name: lambda text, name=name: re.sub(rf"\$\`UNIT\.{re.escape(name)}\`|\$\`\.?{re.escape(name)}\`", str(item[name]), text) for name in item.keys()}
    values.update(PRESET_VARIABLES)
    values["current_date_formatted"] = substitute_with_current_date
    
    for paragraph in document.paragraphs:
        if paragraph.text.__contains__("$`"):
            for variable_name, replacer_function in values.items():
                paragraph.text = replacer_function(paragraph.text)
    document.save(output_path)
    try:
        with open("./contents/config.txt") as f:
            if any([line.split(' ')[0] == "open_file_after_generation" and line.split(' ')[1].strip().lower() == "true" for line in f.readlines()]):
                open_word_document(output_path)
    except:None
    
@Slot()
def get_item_list_from_csv(csv_path: str) -> list[dict[str, str]]:
    with open(csv_path, 'r', encoding='utf-8') as f:
        item_list = [{k: v for k, v in row.items()} for row in csv.DictReader(f, skipinitialspace=True)]
    return item_list

@Slot()
def expect_items_from_docx(template_path: str) -> dict[str, str | list[str]]:
    document = Document(template_path)
    text = "\n".join([paragraph.text for paragraph in document.paragraphs])
    unit_variables = list(set([re.sub(rf"\$|\`|UNIT|\.","",var) for var in (re.findall(rf"\$\`UNIT\.\w+\`|\$\`\.?\w+\`", text))]))
    item_variables = [re.sub(rf"\$|\`","",var) for var in list(set(re.findall(rf"\$`\w+[.]\w+`", text)))]

    items = dict()
    for item_name in list(set([var.split('.')[0] for var in item_variables])):
        items[item_name] = list(set([var.split('.')[1] for var in item_variables if var.split('.')[0]==item_name]))
    items["UNIT"] = unit_variables
    for item_var_lists in items.values(): item_var_lists.sort()
    
    return items

@Slot()
def generate_final_input_item(expected_attributes:dict[str, str | list[str]], chosen_attributes:dict[str, str | list[str]]) -> dict[str, str]:
    final = dict()
    
    for item_name in [item_name for item_name in expected_attributes.keys() if item_name != "UNIT"]:
        item = [item for item in get_item_list_from_csv(chosen_attributes[item_name]['path']) if item['id']==chosen_attributes[item_name]['id']][0]
        for var in item.keys():
            final[f"{item_name}.{var}"] = item[var]
            
    for unit_name in expected_attributes['UNIT']:
        if not PRESET_VARIABLES.__contains__(unit_name):
            final[unit_name] = chosen_attributes[unit_name]

    return final


@Slot()
def find_fitting_items(expected_attributes : list[str]) -> dict[str, str]:
    PATH = "./contents/items"
    item_files = os.listdir(PATH)
    fitting_items = dict()
    for item_file in item_files:
        items = get_item_list_from_csv(f"{PATH}/{item_file}")
        if all(items[0].__contains__(x) for x in expected_attributes):
            fitting_items.update({item["id"]:item_file for item in items})
                
    return fitting_items